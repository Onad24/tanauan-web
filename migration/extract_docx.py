#!/usr/bin/env python3
"""Extract text from all DOCX files in the drive_data tree."""
import os, sys
sys.path.insert(0, '/tmp/drivetools/lib/python3.11/site-packages')
from docx import Document

root = '/home/krixan/Documents/tanauan/tanauan-leyte-clone/migration/drive_data'
outdir = '/home/krixan/Documents/tanauan/tanauan-leyte-clone/migration/extracted'
os.makedirs(outdir, exist_ok=True)

for dirpath, _, files in os.walk(root):
    for f in files:
        if f.lower().endswith('.docx'):
            p = os.path.join(dirpath, f)
            try:
                doc = Document(p)
                # Also grab tables
                text_parts = []
                for para in doc.paragraphs:
                    t = para.text.strip()
                    if t:
                        text_parts.append(t)
                for table in doc.tables:
                    for row in table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            text_parts.append(' | '.join(cells))
                text = '\n'.join(text_parts)
                if not text:
                    print(f"EMPTY: {p}")
                    continue
                rel = os.path.relpath(p, root).replace('/', '__').replace(' ', '_')
                out = os.path.join(outdir, rel + '.txt')
                with open(out, 'w') as fh:
                    fh.write(text)
                print(f"OK: {rel} ({len(text)} chars)")
            except Exception as e:
                print(f"ERR: {p}: {e}")